"""
Document parsers: one function per file type.
Each returns List[dict] with keys: text, page_or_row, source_file, (optional) section_title.
"""
from __future__ import annotations

import logging
from typing import Any

log = logging.getLogger(__name__)


def pdf_parser(file_path: str) -> list[dict[str, Any]]:
    """
    Parse a PDF with PyMuPDF (fitz). Returns one segment per page.
    Falls back to pdfplumber if fitz yields an empty page.
    """
    import fitz  # PyMuPDF

    segments: list[dict[str, Any]] = []
    try:
        with fitz.open(file_path) as doc:
            for page_num, page in enumerate(doc, start=1):
                text = page.get_text("text").strip()
                if not text:
                    # Fallback: try pdfplumber for this page
                    text = _pdfplumber_page(file_path, page_num - 1)
                if text:
                    segments.append({
                        "text": text,
                        "page_or_row": str(page_num),
                        "source_file": file_path,
                    })
    except Exception as exc:
        log.error("PDF parsing failed for %s: %s", file_path, exc)
        raise
    return segments


def _pdfplumber_page(file_path: str, page_index: int) -> str:
    """Extract text from a single page via pdfplumber (used as fallback)."""
    try:
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            if page_index < len(pdf.pages):
                return (pdf.pages[page_index].extract_text() or "").strip()
    except Exception as exc:
        log.debug("pdfplumber fallback failed for %s page %d: %s", file_path, page_index, exc)
    return ""


def excel_parser(file_path: str) -> list[dict[str, Any]]:
    """
    Parse an Excel workbook. Returns one text segment per data row.
    Row text is formatted as "col: value | col: value …" for dense embedding context.
    """
    import openpyxl

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    segments: list[dict[str, Any]] = []
    headers: list[str] = []

    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_idx == 1:
            headers = [
                str(cell).strip() if cell is not None else f"col_{i}"
                for i, cell in enumerate(row)
            ]
            continue
        row_dict = {h: v for h, v in zip(headers, row)}
        parts = [f"{k}: {v}" for k, v in row_dict.items() if v is not None]
        if parts:
            segments.append({
                "text": " | ".join(parts),
                "page_or_row": str(row_idx),
                "source_file": file_path,
                "row_data": row_dict,
            })

    wb.close()
    return segments


def ocr_parser(file_path: str) -> list[dict[str, Any]]:
    """
    Extract text from a scanned image using pytesseract.
    PSM 6 treats the image as a uniform block of text — best for printed reports.
    """
    import pytesseract
    from PIL import Image

    img = Image.open(file_path).convert("L")  # grayscale improves OCR accuracy
    text = pytesseract.image_to_string(img, config="--psm 6 --oem 3").strip()
    if not text:
        log.warning("OCR produced no text for %s", file_path)
        return []
    return [{
        "text": text,
        "page_or_row": "1",
        "source_file": file_path,
    }]


def docx_parser(file_path: str) -> list[dict[str, Any]]:
    """
    Parse a DOCX file. Splits on Heading styles to produce section-level segments.
    Each segment carries its section_title for downstream chunking awareness.
    """
    from docx import Document

    doc = Document(file_path)
    segments: list[dict[str, Any]] = []
    buffer: list[str] = []
    current_heading: str | None = None

    def _flush() -> None:
        if buffer:
            segments.append({
                "text": "\n".join(buffer).strip(),
                "page_or_row": current_heading or "body",
                "source_file": file_path,
                "section_title": current_heading,
            })
            buffer.clear()

    for para in doc.paragraphs:
        if para.style.name.startswith("Heading"):
            _flush()
            current_heading = para.text.strip()
        elif para.text.strip():
            buffer.append(para.text.strip())

    _flush()
    return segments
